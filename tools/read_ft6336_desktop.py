import openpyxl
import os

path = r'd:\360MoveData\Users\Administrator\Desktop\触摸驱动\Ft6336GU_Firmware 外部寄存器_20151112(2).xlsx'
outpath = r'e:\github\github_code\all\4.8code\tools\ft6336_regs_desktop.txt'

with open(outpath, 'w', encoding='utf-8') as out:
    out.write(f"File exists: {os.path.exists(path)}\n")
    out.write(f"File size: {os.path.getsize(path) if os.path.exists(path) else 'N/A'}\n\n")
    
    wb = openpyxl.load_workbook(path, data_only=True)
    out.write(f"Sheets: {wb.sheetnames}\n")
    for sn in wb.sheetnames:
        ws = wb[sn]
        out.write(f"\n=== {sn} (rows={ws.max_row}, cols={ws.max_column}) ===\n")
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i > 300:
                out.write("... (truncated)\n")
                break
            vals = [str(c) if c is not None else "" for c in row]
            out.write(" | ".join(vals) + "\n")

print("Done")
